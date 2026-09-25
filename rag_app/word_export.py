"""Editable Word reports with native headings, lists and internal navigation."""

from io import BytesIO
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def clean(text):
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(text))


def word_report(task):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.2)
    section.left_margin = section.right_margin = Cm(2.4)
    for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'List Bullet']:
        style = doc.styles[name]
        style.font.name = 'Calibri'
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        style.paragraph_format.space_after = Pt(8)
    normal = doc.styles['Normal']
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.first_line_indent = Pt(22)
    doc.styles['Title'].font.size = Pt(24)
    doc.styles['Heading 1'].font.size = Pt(16)
    for name in ['Title', 'Heading 1', 'Heading 2', 'List Bullet']:
        doc.styles[name].paragraph_format.first_line_indent = Pt(0)
    doc.core_properties.title = clean(task['title'])
    doc.core_properties.author = '知识资料工作台'
    doc.add_paragraph(clean(task['title']), 'Title')
    if task.get('summary'):
        doc.add_paragraph('汇总答案', 'Heading 1')
        for line in clean(task['summary']).splitlines():
            if line.strip():
                doc.add_paragraph(line[2:] if line.startswith('• ') else line,
                                  'List Bullet' if line.startswith('• ') else 'Normal')
    doc.add_paragraph('本资料按任务清单整理知识库中的相关内容。正文可以直接编辑；使用下方章节链接或 Word 导航窗格可跳转到对应部分。')
    if task['status'] != 'completed':
        doc.add_paragraph('本资料尚未全部完成，部分内容需要补充或重试。')
    doc.add_paragraph('章节导航', 'Heading 2')
    for index, step in enumerate(task['steps'], 1):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.first_line_indent = Pt(0)
        link = OxmlElement('w:hyperlink')
        link.set(qn('w:anchor'), f'section_{index}')
        run = OxmlElement('w:r')
        props = OxmlElement('w:rPr')
        color = OxmlElement('w:color')
        color.set(qn('w:val'), '2257C7')
        props.append(color)
        run.append(props)
        text = OxmlElement('w:t')
        text.text = clean(step['topic'])
        run.append(text)
        link.append(run)
        paragraph._p.append(link)
    for index, step in enumerate(task['steps'], 1):
        heading = doc.add_paragraph(clean(step['topic']), 'Heading 1')
        start = OxmlElement('w:bookmarkStart')
        start.set(qn('w:id'), str(index))
        start.set(qn('w:name'), f'section_{index}')
        end = OxmlElement('w:bookmarkEnd')
        end.set(qn('w:id'), str(index))
        heading._p.append(start)
        heading._p.append(end)
        for line in clean(step.get('answer') or '此部分尚未完成。').splitlines():
            line = line.strip()
            if line:
                doc.add_paragraph(line[2:] if line.startswith('• ') else line,
                                  'List Bullet' if line.startswith('• ') else 'Normal')
    footer = section.footer.paragraphs[0]
    footer.paragraph_format.first_line_indent = Pt(0)
    footer.alignment = 2
    footer.add_run('第 ')
    page = OxmlElement('w:fldSimple')
    page.set(qn('w:instr'), 'PAGE')
    footer._p.append(page)
    footer.add_run(' 页')
    output = BytesIO()
    doc.save(output)
    return output.getvalue()
